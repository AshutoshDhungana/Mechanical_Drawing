"""Extract lookup tables and a golden snapshot from Vertical-Lift-Gate.xlsm.

Usage:  python tools/extract_workbook.py [path/to/workbook.xlsm]
Writes: data/tables/*.json   lookup tables, source units, with cell provenance
        data/golden/workbook.json   every non-empty cell: formula + cached value
        data/golden/inputs.json     editable (unlocked) cells with label/unit/validation
Needs:  pip install openpyxl
"""
import json, re, sys, warnings, datetime
from pathlib import Path
import openpyxl
from openpyxl.utils import range_boundaries, get_column_letter

warnings.filterwarnings("ignore")  # openpyxl: unsupported header/CF/DV extensions
ROOT = Path(__file__).resolve().parent.parent
SRC = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "Vertical-Lift-Gate.xlsm"
OUT = ROOT / "data"

wf = openpyxl.load_workbook(SRC)                   # formulas
wv = openpyxl.load_workbook(SRC, data_only=True)   # values cached by Excel at last save


def jsonable(v):
    if isinstance(v, (datetime.datetime, datetime.date, datetime.time)):
        return v.isoformat()
    return v


def num(v):
    """Numbers stored as padded text in some sheets (e.g. ' 410   ')."""
    if isinstance(v, str):
        try:
            return float(v.strip())
        except ValueError:
            return v
    return v


def rows(sheet, rng):
    return [[num(c.value) for c in r] for r in wv[sheet][rng]]


def table(sheet, rng, cols, units, note=None, skip_blank=True):
    data = [dict(zip(cols, r)) for r in rows(sheet, rng)]
    if skip_blank:
        data = [d for d in data if d[cols[0]] not in (None, "")]
    t = {"source": f"'{sheet}'!{rng}", "units": units, "rows": data}
    if note:
        t["note"] = note
    return t


def col(sheet, rng):
    return [v for r in rows(sheet, rng) for v in r if v not in (None, "")]


def split_value_label(items):
    """'0.02 antifriction bearings (static)' -> {value: 0.02, label: ...}"""
    out = []
    for s in items:
        m = re.match(r"\s*([\d.]+)\s*(.*)", str(s))
        out.append({"value": float(m[1]), "label": m[2]} if m else {"value": None, "label": s})
    return out


# ---------------------------------------------------------------- tables
SEC = ["designation", "w", "area", "h", "b", "tf", "tw"]
SEC_U = {"w": "kg/m", "area": "cm2", "h": "mm", "b": "mm", "tf": "mm", "tw": "mm"}
tables = {}

tables["ismb"] = table(
    "ISMB", "A6:M19", SEC + ["Ixx", "Iyy", "rxx", "ryy", "Zxx", "Zyy"],
    {**SEC_U, "Ixx": "cm4", "Iyy": "cm4", "rxx": "cm", "ryy": "cm", "Zxx": "cm3", "Zyy": "cm3"},
    "IS 808 medium-weight beams.")

tables["ismc"] = table(
    "ISMC", "A5:U15",
    SEC + ["Cyy", "Ixx", "Iyy", "rxx", "ryy", "Zxx", "Zyy", "h1", "h2", "b1_half", "C", "g", "g1_min", "max_rivet"],
    {**SEC_U, "Cyy": "cm", "Ixx": "cm4", "Iyy": "cm4", "rxx": "cm", "ryy": "cm", "Zxx": "cm3", "Zyy": "cm3",
     "h1": "mm", "h2": "mm", "b1_half": "mm", "C": "mm", "g": "mm", "g1_min": "mm", "max_rivet": "mm"},
    "IS 808 medium-weight channels; h1..max_rivet are the sheet's 'Connection Details'.")

tables["gate_shapes"] = table(
    "Lib", "A5:W18",
    ["designation", "back_flange_width", "back_flange_t", "web_t", "web_depth", "flange_width", "flange_t",
     "Ixx", "weight"] + [None] * 12 + ["area", "Zx"],
    {"back_flange_width": "m", "back_flange_t": "m", "web_t": "m", "web_depth": "m", "flange_width": "m",
     "flange_t": "m", "Ixx": "cm4", "weight": "kg/m", "area": "m2", "Zx": "m3"},
    "Named range ShapeData (A5:G18). ISMB re-expressed as a 3-plate shape: web_depth = h - 2*tf. "
    "Columns J..U hold unrelated dropdown lists (see lists.json) and are dropped.")
for r in tables["gate_shapes"]["rows"]:
    r.pop(None, None)

tables["ismc_lib"] = table(
    "Lib", "A22:I32", ["designation", "w", "area", "h", "b", "tf", "tw", "Ixx", "Zxx"],
    {"w": "kg/m", "area": "cm2", "h": "mm", "b": "mm", "tf": "mm", "tw": "mm", "Ixx": "cm4", "Zxx": "cm3"},
    "ISMC subset used only by the hidden 'Embedded' sheet; duplicate of ismc.json.")

ROPE = ["d", "mass_CF", "mass_CWR", "mbf_1570_CF", "mbf_1570_CWR", "mbf_1770_CF", "mbf_1770_CWR",
        "mbf_1960_CF", "mbf_1960_CWR"]
tables["rope_6x37"] = table(
    "6x37 Rope specification", "B7:J31", ROPE,
    {"d": "mm", "mass_CF": "kg/m", "mass_CWR": "kg/m", **{k: "kN" for k in ROPE[3:]}},
    "6x37 M (18/12/6-1) wire rope, minimum breaking force by grade (N/mm2) and core. "
    "0 = size not made with that core. Used by 'Hoist(Rope Drum)'.")
tables["rope_6x37_old"] = table(
    "Sheet3", "A10:I34", ROPE,
    {"d": "mm", "mass_CF": "kg/100m", "mass_CWR": "kg/100m", **{k: "kN" for k in ROPE[3:]}},
    "Named range RopeData; used only by the hidden legacy 'Rope Drum Hoist'. Same data as rope_6x37 "
    "but mass per 100 m - see data/README.md for the typos it contains.")
# Typos in the active rope sheet, found by cross-checking Sheet3 and the d^2 trend. Rows above stay
# exactly as in the workbook (golden tests need them); the app applies these corrections.
tables["rope_6x37"]["errata"] = [
    {"d": 52, "field": "mbf_1770_CWR", "workbook": 11524, "corrected": 1524},
    {"d": 64, "field": "mass_CF", "workbook": 14.71, "corrected": 14.17},
]

tables["steel_grades"] = table(
    "material ", "B5:F13", ["grade", "uts", "fy_t_lt20", "fy_t_20_40", "fy_t_gt40"],
    {"uts": "MPa", "fy_t_lt20": "MPa", "fy_t_20_40": "MPa", "fy_t_gt40": "MPa"},
    "IS 2062 grades; yield stress by plate thickness band (<20, 20-40, >40 mm). Not referenced by any formula.")

# Anchor bolts: the sheet's own size labels are an Excel fill-series accident (M24,M25,M26,M27 ...);
# the lookup formulas address rows by block, so the block (not the label) is the real size.
bolt_blocks = {"M16": (71, 73), "M20": (75, 77), "M24": (79, 82), "M27": (83, 86)}
tables["anchor_bolts"] = {
    "source": "'Embedded Parts'!AA70:AC86",
    "units": {"length": "mm", "col_AC": "unlabelled (probably kg per bolt)"},
    "note": "Size taken from the row block the formulas at AA88:AB95 address, not from column AA labels.",
    "rows": [{"size": size, "length": wv["Embedded Parts"][f"AB{r}"].value,
              "col_AC": wv["Embedded Parts"][f"AC{r}"].value, "sheet_label": wv["Embedded Parts"][f"AA{r}"].value}
             for size, (a, b) in bolt_blocks.items() for r in range(a, b + 1)],
}


def plate_table(name, b_rng, raw_rng, used_rng, heads):
    ratio = col("Plates", b_rng)
    raw, used = rows("Plates", raw_rng), rows("Plates", used_rng)
    return {
        "source": f"'Plates'!{b_rng} (b/a), raw {raw_rng}, as-used {used_rng}",
        "note": f"{name}. Bending moment M = coef * p * a^2. 'raw' are the textbook coefficients; the sheet "
                "multiplies them by 0.0254 ('as_used') and Stress divides by 0.0254 again, so the factor "
                "cancels. Lookup in Stress is linear interpolation between bracketing b/a rows "
                "(FORECAST over OFFSET(MATCH(...,1))).",
        "b_over_a": ratio,
        "columns": [{"name": n, "load": ld, "moment": m, "point": pt, "location": loc,
                     "raw": [r[i] for r in raw[:len(ratio)]], "as_used": [r[i] for r in used[:len(ratio)]]}
                    for i, (n, ld, m, pt, loc) in enumerate(heads)],
    }


tables["plate_coefficients"] = {
    "CCCF": plate_table(
        "Edge (top/bottom) beam panels: three edges fixed, one free", "B33:B56", "K33:P56", "C33:H56",
        [("CCCFA", "uniform", "Mx", "4,6", "x=0,y=0"), ("CCCFB", "uniform", "My", "1", "x=b/2,y=a"),
         ("CCCFD", "hydrostatic A", "Mx", "4", "x=0,y=0"), ("CCCFE", "hydrostatic A", "My", "1", "x=b/2,y=a"),
         ("CCCFG", "hydrostatic C", "Mx", "6", "x=0,y=0"), ("CCCFH", "hydrostatic C", "My", "1", "x=b/2,y=a")]),
    "CCCC": plate_table(
        "Intermediate beam panels: four edges fixed", "B89:B111", "P89:X111", "C89:K111",
        [("CCCCA", "uniform", "Mx", "4,6", "x=0,y=0"), ("CCCCB", "uniform", "Mx", "2", "x=0,y=a/2"),
         ("CCCCC", "uniform", "My", "1", "x=b/2,y=a/2"), ("CCCCD", "uniform", "My", "2", "x=0,y=a/2"),
         ("CCCCE", "hydrostatic", "Mx", "4", "x=0,y=0"), ("CCCCF", "hydrostatic", "Mx", "6", "x=0,y=a"),
         ("CCCCG", "hydrostatic", "Mx", "2", "x=0,y=a/2"), ("CCCCH", "hydrostatic", "My", "1", "x=b/2,y=a/2"),
         ("CCCCI", "hydrostatic", "My", "2", "x=0,y=a/2")]),
}

tables["lists"] = {
    "beam_types": {"source": "Lib!J5:J7 (BeamTypes)", "values": col("Lib", "J5:J7")},
    "gate_types": {"source": "Lib!N5:N7 (GateType)", "values": col("Lib", "N5:N7")},
    "corrosion_allowance_mm": {"source": "Lib!N10:N20 (CorAll)", "values": col("Lib", "N10:N20")},
    "friction": {"source": "Lib!Q5:Q10 (friction)", "values": split_value_label(col("Lib", "Q5:Q10"))},
    "ballast_density_kg_m3": {"source": "Lib!Q15:Q16 (Ballast)", "values": split_value_label(col("Lib", "Q15:Q16"))},
    "bolt_sizes": {"source": "Lib!Q22:Q30 (CType)", "values": col("Lib", "Q22:Q30")},
    "beam_spacing_methods": {"source": "'Beam Spacing'!N6:N9", "values": col("Beam Spacing", "N6:N9")},
    "beam_spacing_rounding_m": {"source": "'Beam Spacing'!Q5:Q10 (0, 1/16..1 inch in m)",
                                "values": col("Beam Spacing", "Q5:Q10")},
    "wheel_spacing_methods": {"source": "'Wheel Spacing'!K6:K7", "values": col("Wheel Spacing", "K6:K7")},
    "gear_modules_mm": {"source": "'Hoist(Rope Drum)'!W3:W19", "values": col("Hoist(Rope Drum)", "W3:W19")},
    "motor_ratings_kW": {"source": "Sheet3!L10:L29", "values": col("Sheet3", "L10:L29")},
}

# ---------------------------------------------------------------- golden snapshot
dv_by_sheet = {}
for ws in wf.worksheets:
    m = {}
    for d in ws.data_validations.dataValidation:
        for rg in str(d.sqref).split():
            c1, r1, c2, r2 = range_boundaries(rg)
            for r in range(r1, r2 + 1):
                for c in range(c1, c2 + 1):
                    m[f"{get_column_letter(c)}{r}"] = {"type": d.type, "formula1": d.formula1,
                                                       "formula2": d.formula2, "operator": d.operator}
    dv_by_sheet[ws.title] = m

snapshot = {
    "source": SRC.name,
    "note": "Values are what Excel cached at last save (2023-08-18); errors kept as strings like '#DIV/0!'.",
    "names": {d.name: d.attr_text for d in wf.defined_names.values()},  # 18 sheet-local print areas omitted
    "sheets": {},
}
inputs = {"source": SRC.name, "sheets": {}}

for ws in wf.worksheets:
    vs = wv[ws.title]
    cells, ins = {}, []
    for row in ws.iter_rows():
        for c in row:
            v = vs[c.coordinate].value
            f = c.value if isinstance(c.value, str) and c.value.startswith("=") else None
            if f is None and v is None:
                continue
            e = {"v": jsonable(v)}
            if f:
                e["f"] = f
            if c.number_format != "General":
                e["fmt"] = c.number_format
            unlocked = c.protection.locked is False
            if unlocked:
                e["in"] = True
            cells[c.coordinate] = e
            if unlocked and (v is not None or f):
                label = " ".join(str(ws.cell(c.row, k).value) for k in range(1, c.column)
                                 if ws.cell(c.row, k).value is not None
                                 and not str(ws.cell(c.row, k).value).startswith("="))
                unit = re.search(r'"([^"]+)"', c.number_format)  # units live in formats like 0.000 "m"
                unit = unit[1] if unit else vs.cell(c.row, c.column + 1).value
                if isinstance(unit, str) and unit.startswith("ISM"):  # next cell is a section, not a unit
                    unit = None
                ins.append({"cell": c.coordinate, "label": label.strip()[-80:], "value": jsonable(v),
                            "unit": unit if isinstance(unit, str) else None,
                            **({"formula": f} if f else {}),
                            **({"validation": dv_by_sheet[ws.title][c.coordinate]}
                               if c.coordinate in dv_by_sheet[ws.title] else {})})
    snapshot["sheets"][ws.title] = {"state": ws.sheet_state, "dimensions": ws.dimensions, "cells": cells}
    if ins:
        inputs["sheets"][ws.title] = ins

# ---------------------------------------------------------------- write
(OUT / "tables").mkdir(parents=True, exist_ok=True)
(OUT / "golden").mkdir(parents=True, exist_ok=True)
for k, t in tables.items():
    (OUT / "tables" / f"{k}.json").write_text(json.dumps(t, indent=1, ensure_ascii=False), encoding="utf-8")
(OUT / "golden" / "workbook.json").write_text(json.dumps(snapshot, ensure_ascii=False), encoding="utf-8")
(OUT / "golden" / "inputs.json").write_text(json.dumps(inputs, indent=1, ensure_ascii=False), encoding="utf-8")

n_cells = sum(len(s["cells"]) for s in snapshot["sheets"].values())
n_f = sum(1 for s in snapshot["sheets"].values() for c in s["cells"].values() if "f" in c)
n_in = sum(len(v) for v in inputs["sheets"].values())
print(f"tables: {', '.join(tables)}")
print(f"snapshot: {len(snapshot['sheets'])} sheets, {n_cells} cells, {n_f} formulas, {n_in} inputs, "
      f"{len(snapshot['names'])} names")

# ---------------------------------------------------------------- cross-checks (report, don't fail)
issues = []
ismb = {r["designation"]: r for r in tables["ismb"]["rows"]}
for g in tables["gate_shapes"]["rows"]:
    s = ismb[g["designation"]]
    exp = {"flange_width": s["b"] / 1000, "flange_t": s["tf"] / 1000, "web_t": s["tw"] / 1000,
           "web_depth": (s["h"] - 2 * s["tf"]) / 1000, "Ixx": s["Ixx"], "weight": s["w"]}
    for k, e in exp.items():
        if abs(g[k] - e) > 1e-6 * max(1, abs(e)):
            issues.append(f"gate_shapes {g['designation']}.{k}={g[k]} vs ISMB {e}")
new = {r["d"]: r for r in tables["rope_6x37"]["rows"]}
for o in tables["rope_6x37_old"]["rows"]:
    n = new.get(o["d"])
    if not n:
        issues.append(f"rope d={o['d']} only in Sheet3"); continue
    for k in ROPE[1:]:
        ov = o[k] / 100 if k.startswith("mass") else o[k]
        if abs((n[k] or 0) - (ov or 0)) > 1e-6:
            issues.append(f"rope d={o['d']} {k}: rope sheet {n[k]} vs Sheet3 {o[k]}")
for t in tables["plate_coefficients"].values():
    for c in t["columns"]:
        for i, (r, u) in enumerate(zip(c["raw"], c["as_used"])):
            if abs(r * 0.0254 - u) > 1e-12:
                issues.append(f"plate {c['name']} b/a={t['b_over_a'][i]}: raw*0.0254={r * 0.0254} vs {u}")
errs = {}
for name, s in snapshot["sheets"].items():
    n = sum(1 for c in s["cells"].values() if isinstance(c["v"], str) and c["v"].startswith("#"))
    if n:
        errs[name] = n
print("cached error cells:", errs)
print("cross-check issues:" if issues else "cross-checks: all consistent")
for i in issues:
    print("  -", i)
